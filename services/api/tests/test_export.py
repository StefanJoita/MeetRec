# services/api/tests/test_export.py
# ============================================================
# Teste pentru fix-ul #1: Export fără verificare RBAC
# ============================================================
# Verificăm că endpoint-ul GET /export/recording/{id}
# aplică corect check_recording_access() pentru participanți.
#
# Scenarii:
#   1. Participant fără acces → 403 Forbidden
#   2. Participant cu acces   → 200 OK (export livrat)
#   3. Admin                  → 200 OK (acces total)
#   4. Operator               → 200 OK (acces total)
# ============================================================

import uuid
from datetime import date, datetime, timezone
from unittest.mock import AsyncMock, MagicMock, patch, PropertyMock

import pytest
import pytest_asyncio
from httpx import AsyncClient, ASGITransport

from src.main import app
from src.database import get_db
from src.middleware.auth import get_current_user
from src.models.audit_log import User


# ── Helpers ──────────────────────────────────────────────────

def make_user(role: str) -> MagicMock:
    """Creează un User mock cu rolul specificat."""
    user = MagicMock(spec=User)
    user.id = uuid.uuid4()
    user.username = f"user_{role}"
    user.is_active = True
    user.is_admin = role == "admin"
    user.is_participant = role == "participant"
    user.must_change_password = False
    return user


def make_recording_mock(recording_id: uuid.UUID) -> MagicMock:
    rec = MagicMock()
    rec.id = recording_id
    rec.title = "Ședință test"
    rec.meeting_date = date(2024, 3, 15)
    rec.duration_formatted = "00:05:00"
    rec.classification = "N"
    return rec


def make_transcript_mock() -> MagicMock:
    seg1 = MagicMock()
    seg1.start_time = 0.0
    seg1.text = "Bună ziua, deschidem ședința."

    transcript = MagicMock()
    transcript.status = "completed"
    transcript.language = "ro"
    transcript.word_count = 5
    transcript.segments = [seg1]
    return transcript


@pytest_asyncio.fixture
async def export_client(mock_db):
    """Client cu DB mock pentru testele de export."""
    async def _mock_get_db():
        yield mock_db

    app.dependency_overrides[get_db] = _mock_get_db
    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://test",
    ) as ac:
        yield ac
    app.dependency_overrides.clear()


# ── Teste ─────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_export_participant_fara_acces_primeste_403(export_client, mock_db):
    """
    Un participant care NU e linkat la înregistrare
    trebuie să primească 403, nu conținutul exportului.
    """
    recording_id = uuid.uuid4()
    participant = make_user("participant")

    app.dependency_overrides[get_current_user] = lambda: participant

    with patch(
        "src.routers.export._get_transcript_and_recording",
        new_callable=AsyncMock,
        return_value=(make_recording_mock(recording_id), make_transcript_mock()),
    ), patch(
        "src.routers.export.check_recording_access",
        new_callable=AsyncMock,
        return_value=False,   # ← participant fără acces
    ), patch(
        "src.routers.export.log_audit",
        new_callable=AsyncMock,
    ):
        resp = await export_client.get(f"/api/v1/export/recording/{recording_id}?format=txt")

    assert resp.status_code == 403
    assert "interzis" in resp.json()["detail"].lower()


@pytest.mark.asyncio
async def test_export_participant_cu_acces_primeste_200(export_client, mock_db):
    """
    Un participant linkat explicit la înregistrare
    trebuie să poată exporta (200 + conținut txt).
    """
    recording_id = uuid.uuid4()
    participant = make_user("participant")

    app.dependency_overrides[get_current_user] = lambda: participant

    with patch(
        "src.routers.export._get_transcript_and_recording",
        new_callable=AsyncMock,
        return_value=(make_recording_mock(recording_id), make_transcript_mock()),
    ), patch(
        "src.routers.export.check_recording_access",
        new_callable=AsyncMock,
        return_value=True,    # ← participant cu acces
    ), patch(
        "src.routers.export.log_audit",
        new_callable=AsyncMock,
    ):
        resp = await export_client.get(f"/api/v1/export/recording/{recording_id}?format=txt")

    assert resp.status_code == 200
    assert b"Sedinta test" in resp.content or b"test" in resp.content.lower()


@pytest.mark.asyncio
async def test_export_admin_primeste_200_fara_check_participanti(export_client, mock_db):
    """
    Adminul are acces la orice înregistrare.
    check_recording_access returnează True direct pentru non-participanți.
    """
    recording_id = uuid.uuid4()
    admin = make_user("admin")

    app.dependency_overrides[get_current_user] = lambda: admin

    with patch(
        "src.routers.export._get_transcript_and_recording",
        new_callable=AsyncMock,
        return_value=(make_recording_mock(recording_id), make_transcript_mock()),
    ), patch(
        "src.routers.export.check_recording_access",
        new_callable=AsyncMock,
        return_value=True,
    ), patch(
        "src.routers.export.log_audit",
        new_callable=AsyncMock,
    ):
        resp = await export_client.get(f"/api/v1/export/recording/{recording_id}?format=txt")

    assert resp.status_code == 200


@pytest.mark.asyncio
async def test_export_operator_primeste_200(export_client, mock_db):
    """Operatorul are acces la toate înregistrările."""
    recording_id = uuid.uuid4()
    operator = make_user("operator")

    app.dependency_overrides[get_current_user] = lambda: operator

    with patch(
        "src.routers.export._get_transcript_and_recording",
        new_callable=AsyncMock,
        return_value=(make_recording_mock(recording_id), make_transcript_mock()),
    ), patch(
        "src.routers.export.check_recording_access",
        new_callable=AsyncMock,
        return_value=True,
    ), patch(
        "src.routers.export.log_audit",
        new_callable=AsyncMock,
    ):
        resp = await export_client.get(f"/api/v1/export/recording/{recording_id}?format=txt")

    assert resp.status_code == 200


@pytest.mark.asyncio
async def test_export_fara_autentificare_primeste_401(export_client):
    """Fără token JWT → 401."""
    recording_id = uuid.uuid4()
    # Fără override pe get_current_user → middleware-ul respinge
    resp = await export_client.get(f"/api/v1/export/recording/{recording_id}?format=txt")
    assert resp.status_code == 401


# ── Diarizare în export ───────────────────────────────────────

def make_segment(start: float, text: str, speaker_id):
    seg = MagicMock()
    seg.start_time = start
    seg.text = text
    seg.speaker_id = speaker_id
    return seg


def make_diarized_transcript() -> MagicMock:
    transcript = make_transcript_mock()
    transcript.segments = [
        make_segment(0.0, "Bună ziua, deschidem ședința.", "SPEAKER_00"),
        make_segment(4.0, "Mulțumesc, încep eu.", "SPEAKER_01"),
        make_segment(9.0, "Fără vorbitor.", None),
    ]
    return transcript


async def export_with_speakers(export_client, mock_db, fmt: str, speaker_mapping: dict, users=()):
    recording_id = uuid.uuid4()
    recording = make_recording_mock(recording_id)
    recording.speaker_mapping = speaker_mapping
    app.dependency_overrides[get_current_user] = lambda: make_user("operator")

    users_result = MagicMock()                       # SELECT users pentru speaker_mapping
    users_result.scalars.return_value.all.return_value = list(users)
    mock_db.execute = AsyncMock(return_value=users_result)
    with patch(
        "src.routers.export._get_transcript_and_recording",
        new_callable=AsyncMock,
        return_value=(recording, make_diarized_transcript()),
    ), patch(
        "src.routers.export.check_recording_access", new_callable=AsyncMock, return_value=True,
    ), patch(
        "src.routers.export.log_audit", new_callable=AsyncMock,
    ):
        return await export_client.get(f"/api/v1/export/recording/{recording_id}?format={fmt}")


def docx_text(content: bytes) -> str:
    import io
    from docx import Document
    return "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)


@pytest.mark.asyncio
async def test_export_txt_shows_unmapped_speakers(export_client, mock_db):
    """Vorbitorii neasociați cu un utilizator apar ca „Vorbitor N”."""
    resp = await export_with_speakers(export_client, mock_db, "txt", {})

    text = resp.content.decode("utf-8")
    assert resp.status_code == 200
    assert "[00:00] Vorbitor 1: Bună ziua, deschidem ședința." in text
    assert "[00:04] Vorbitor 2: Mulțumesc, încep eu." in text
    assert "[00:09]  Fără vorbitor." in text


@pytest.mark.asyncio
async def test_export_docx_mixes_mapped_and_unmapped_speakers(export_client, mock_db):
    user_id = uuid.uuid4()
    user = MagicMock(id=user_id, full_name="Ion Popescu", username="ipopescu")

    resp = await export_with_speakers(export_client, mock_db, "docx", {"SPEAKER_00": str(user_id)}, [user])

    text = docx_text(resp.content)
    assert resp.status_code == 200
    assert "[00:00] Ion Popescu" in text
    assert "[00:04] Vorbitor 2" in text


@pytest.mark.asyncio
async def test_export_speaker_mapped_to_deleted_user(export_client, mock_db):
    """Utilizatorul asociat a fost șters → eticheta implicită, nu dispare vorbitorul."""
    resp = await export_with_speakers(export_client, mock_db, "txt", {"SPEAKER_01": str(uuid.uuid4())})

    assert "Vorbitor 2: Mulțumesc, încep eu." in resp.content.decode("utf-8")


@pytest.mark.asyncio
async def test_export_pdf_with_speakers(export_client, mock_db):
    resp = await export_with_speakers(export_client, mock_db, "pdf", {})

    assert resp.status_code == 200
    assert resp.content.startswith(b"%PDF")


# ── Paragrafe pe vorbitor ─────────────────────────────────────

def test_consecutive_segments_of_same_speaker_become_one_paragraph():
    from src.routers.export import _paragraphs
    segments = [
        make_segment(0.0, "Bună ziua.", "SPEAKER_00"),
        make_segment(3.0, "Începem ședința.", "SPEAKER_00"),
        make_segment(6.0, "Mulțumesc.", "SPEAKER_01"),
        make_segment(8.0, "Fără vorbitor.", None),
        make_segment(9.0, "Tot fără.", None),
        make_segment(10.0, "Revin eu.", "SPEAKER_00"),
    ]
    labels = {"SPEAKER_00": "Ion Popescu", "SPEAKER_01": "Vorbitor 2"}

    assert _paragraphs(segments, labels) == [
        (0.0, "Ion Popescu", "Bună ziua. Începem ședința."),
        (6.0, "Vorbitor 2", "Mulțumesc."),
        (8.0, None, "Fără vorbitor."),                  # fără vorbitor: segmente separate
        (9.0, None, "Tot fără."),
        (10.0, "Ion Popescu", "Revin eu."),
    ]


@pytest.mark.asyncio
async def test_export_uses_free_text_speaker_names(export_client, mock_db):
    resp = await export_with_speakers(export_client, mock_db, "txt", {"SPEAKER_00": "Ion <Popescu> & co"})

    assert "[00:00] Ion <Popescu> & co: Bună ziua, deschidem ședința." in resp.content.decode("utf-8")


@pytest.mark.asyncio
async def test_pdf_survives_markup_characters(export_client, mock_db):
    resp = await export_with_speakers(export_client, mock_db, "pdf", {"SPEAKER_00": "Ion <b>Popescu & co"})

    assert resp.status_code == 200 and resp.content.startswith(b"%PDF")


# ── Numele fișierului și nivelul de clasificare ──────────────

@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["txt", "pdf", "docx"])
async def test_export_nume_fisier_si_nivel_de_clasificare(export_client, mock_db, fmt):
    """Fișierul se numește AAAALLZZ_NIVEL_Titlu.ext; denumirea nivelului apare în document."""
    recording_id = uuid.uuid4()
    app.dependency_overrides[get_current_user] = lambda: make_user("operator")
    recording = make_recording_mock(recording_id)
    recording.classification = "SSv"

    with patch(
        "src.routers.export._get_transcript_and_recording",
        new_callable=AsyncMock,
        return_value=(recording, make_transcript_mock()),
    ), patch(
        "src.routers.export.check_recording_access",
        new_callable=AsyncMock,
        return_value=True,
    ), patch(
        "src.routers.export.log_audit",
        new_callable=AsyncMock,
    ):
        resp = await export_client.get(f"/api/v1/export/recording/{recording_id}?format={fmt}")

    assert resp.status_code == 200
    expected = f"{date.today():%Y%m%d}_SSv_%C8%98edin%C8%9B%C4%83_test.{fmt}"   # Ședință_test
    assert resp.headers["content-disposition"] == f"attachment; filename*=UTF-8''{expected}"
    if fmt == "txt":
        assert resp.text.splitlines()[0].strip() == "SECRET DE SERVICIU"
    elif fmt == "docx":
        import io
        from docx import Document
        header = Document(io.BytesIO(resp.content)).sections[0].header.paragraphs[0].text
        assert header == "SECRET DE SERVICIU"


class TestExportNaming:

    @pytest.mark.parametrize("classification, expected", [
        ("SS", "20261028_SS_Ședința_consiliului.pdf"),
        ("N", "20261028_N_Ședința_consiliului.pdf"),
        (None, "20261028_Ședința_consiliului.pdf"),           # fără înregistrare (jurnalul de audit)
    ])
    def test_format(self, classification, expected):
        from src.services.export_naming import export_filename
        assert export_filename("Ședința consiliului", "pdf", classification, on=date(2026, 10, 28)) == expected

    @pytest.mark.parametrize("title, expected", [
        ('Raport: "buget"/2027?', "Raport_buget2027"),        # caractere interzise în Windows
        ("a\r\nb", "ab"),                                      # fără injecție în header
        ("  multe    spații  ", "multe_spații"),
        ("...", "Document"),
        ("x" * 200, "x" * 80),
    ])
    def test_document_name(self, title, expected):
        from src.services.export_naming import document_name
        assert document_name(title) == expected
